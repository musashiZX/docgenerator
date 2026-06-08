"""Agent-mode Word (.docx) editor powered by Claude Opus 4.7.

Usage:
    python doc_editor.py path/to/document.docx
"""

import json
import os
import sys
from html import escape
from pathlib import Path

from anthropic import Anthropic, APIError, APIConnectionError, RateLimitError
from docx import Document
from dotenv import load_dotenv

load_dotenv()

MODEL = "claude-opus-4-7"

SYSTEM_PROMPT = """You are a Word-document editing agent. The user has opened a \
.docx file and you help edit it through chat. The user may write to you in any \
language (English, Chinese, etc.); the *document language* is whatever the user \
asks for — if unspecified, match the document's existing language.

CORE RULES
1. ALWAYS call `read_document` FIRST at the start of every turn to see the \
current state — paragraph indices, headings, and existing content.
2. When the user asks to add/edit content under a specific section (e.g. \
"add to Education", "在工作经历里加一条", "丰富教育经历"), you MUST:
   a) locate the heading paragraph for that section in the read_document output,
   b) find where that section ends (next heading of equal or higher level, or \
      end of document),
   c) insert the new content INSIDE that section using `insert_paragraph` with \
      the correct index — never just `append_paragraph` to the end of the doc.
3. Paragraph indices are 0-based and shift after every insert/delete. After \
2+ positional edits in a row, call `read_document` again before the next one.
4. Respect the document's structure: a Heading 2 belongs under a Heading 1; \
list items use 'List Bullet' or 'List Number' styles; body text uses 'Normal'.

CONTENT QUALITY
- When the user says "enrich", "expand", "add details", "丰富", "补充", \
"展开", or similar, produce SUBSTANTIVE content — not a single line. For each \
item you add, include the relevant supporting details a reader would expect:
  • Education entry → years, university, department/major, location, plus 1–3 \
    bullets of relevant courses, GPA/honors, thesis topic, or activities.
  • Work entry → years, company, location, role, then 2–4 bullets describing \
    responsibilities and measurable impact.
  • Project entry → name, dates, stack/role, then 2–3 bullets on what was \
    built and the outcome.
- If the user gave you partial facts, fill in plausible, professional placeholder \
bullets and clearly mark anything fabricated, OR ask one concise clarifying \
question before writing — pick whichever is faster for the user.
- Match the existing document's tone, language, and formatting.
- Never leave "[add details here]" placeholders in content you wrote yourself.

TABLES
- `read_document` shows tables as [TABLE t_index] with every row's cell contents.
- To inspect a table in detail use `read_table(table_index)` — it shows every \
  cell as [row][col]: "text".
- To edit a single cell use `set_table_cell(table_index, row, col, text)`.
- To change the same text wherever it appears (including inside tables) use \
  `replace_text` — this is the fastest option when you know the exact string.
- Paragraph tools (set_paragraph, insert_paragraph, etc.) operate on non-table \
  paragraphs only; never use a paragraph index to target a table cell.

OUTPUT
- After editing, give a SHORT (1–3 sentence) summary of what changed AND where \
(which section / table / paragraph indices).
- Available paragraph styles: 'Normal', 'Title', 'Heading 1', 'Heading 2', \
'Heading 3', 'List Bullet', 'List Number', 'Quote'.
"""

TOOLS = [
    {
        "name": "read_document",
        "description": (
            "Read the full document structure in order. "
            "Paragraphs are shown as [p_index] (style) text. "
            "Tables are shown as [TABLE t_index] followed by each row's cell texts. "
            "Use p_index with paragraph tools and t_index with table tools."
        ),
        "input_schema": {"type": "object", "properties": {}},
    },
    {
        "name": "append_paragraph",
        "description": "Append a new paragraph at the end of the document body (outside any table).",
        "input_schema": {
            "type": "object",
            "properties": {
                "text": {"type": "string"},
                "style": {"type": "string", "description": "Word style name (optional)"},
            },
            "required": ["text"],
        },
    },
    {
        "name": "insert_paragraph",
        "description": "Insert a new paragraph before the paragraph at `index` (0-based, outside tables).",
        "input_schema": {
            "type": "object",
            "properties": {
                "index": {"type": "integer"},
                "text": {"type": "string"},
                "style": {"type": "string"},
            },
            "required": ["index", "text"],
        },
    },
    {
        "name": "set_paragraph",
        "description": "Replace the full text (and optionally the style) of the paragraph at `index` (outside tables).",
        "input_schema": {
            "type": "object",
            "properties": {
                "index": {"type": "integer"},
                "text": {"type": "string"},
                "style": {"type": "string"},
            },
            "required": ["index", "text"],
        },
    },
    {
        "name": "delete_paragraph",
        "description": "Delete the paragraph at `index` (outside tables). Later paragraphs shift up by one.",
        "input_schema": {
            "type": "object",
            "properties": {"index": {"type": "integer"}},
            "required": ["index"],
        },
    },
    {
        "name": "replace_text",
        "description": (
            "Find-and-replace a string across the ENTIRE document — "
            "including inside table cells. Returns the number of replacements made. "
            "Prefer this tool when you know the exact text to change."
        ),
        "input_schema": {
            "type": "object",
            "properties": {
                "find": {"type": "string"},
                "replace": {"type": "string"},
            },
            "required": ["find", "replace"],
        },
    },
    {
        "name": "read_table",
        "description": (
            "Return the full contents of a specific table with row and column indices. "
            "Use this before calling set_table_cell to confirm the exact cell coordinates."
        ),
        "input_schema": {
            "type": "object",
            "properties": {
                "table_index": {"type": "integer", "description": "0-based table index from read_document"},
            },
            "required": ["table_index"],
        },
    },
    {
        "name": "set_table_cell",
        "description": (
            "Replace the text of a single table cell identified by table_index, row, and col (all 0-based). "
            "Call read_table first to confirm the correct coordinates."
        ),
        "input_schema": {
            "type": "object",
            "properties": {
                "table_index": {"type": "integer"},
                "row":         {"type": "integer"},
                "col":         {"type": "integer"},
                "text":        {"type": "string"},
            },
            "required": ["table_index", "row", "col", "text"],
        },
    },
]


class DocEditor:
    def __init__(self, path: Path):
        self.path = path
        if path.exists():
            self.doc = Document(str(path))
        else:
            self.doc = Document()
            self.doc.save(str(path))

    def save(self) -> None:
        self.doc.save(str(self.path))

    @staticmethod
    def _clear_runs(paragraph) -> None:
        for run in list(paragraph.runs):
            run._element.getparent().remove(run._element)

    def _apply_style(self, paragraph, style: str | None) -> str | None:
        if not style:
            return None
        try:
            paragraph.style = self.doc.styles[style]
            return None
        except KeyError:
            return f"(warning: style '{style}' not found, using default)"

    # --- helpers for table iteration ---

    def _body_blocks(self):
        """Yield (kind, obj) for each top-level block in document order.

        kind == 'paragraph' → obj is a Paragraph
        kind == 'table'     → obj is a Table
        """
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        for child in self.doc.element.body:
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
            if tag == "p":
                yield "paragraph", Paragraph(child, self.doc)
            elif tag == "tbl":
                yield "table", Table(child, self.doc)

    # --- tool implementations ---

    def read_document(self) -> str:
        """Return full document structure: paragraphs with indices, tables with cell content."""
        lines = []
        p_idx = 0
        t_idx = 0
        has_content = False

        for kind, obj in self._body_blocks():
            has_content = True
            if kind == "paragraph":
                style = obj.style.name if obj.style else "Normal"
                lines.append(f"[{p_idx}] ({style}) {obj.text}")
                p_idx += 1
            else:
                # Table — show every cell so the agent can see and edit all content.
                rows = obj.rows
                n_cols = len(obj.columns) if rows else 0
                lines.append(f"[TABLE {t_idx}] ({len(rows)} rows × {n_cols} cols)")
                for r, row in enumerate(rows):
                    cells = [c.text.replace("\n", " / ") for c in row.cells]
                    lines.append(f"  row {r}: {cells}")
                t_idx += 1

        if not has_content:
            return "(empty document)"
        return "\n".join(lines)

    def append_paragraph(self, text: str, style: str | None = None) -> str:
        p = self.doc.add_paragraph(text)
        warn = self._apply_style(p, style)
        self.save()
        idx = len(self.doc.paragraphs) - 1
        return f"Appended paragraph at index {idx}." + (f" {warn}" if warn else "")

    def insert_paragraph(self, index: int, text: str, style: str | None = None) -> str:
        n = len(self.doc.paragraphs)
        if index < 0 or index > n:
            return f"Error: index {index} out of range (0..{n})."
        if index == n:
            return self.append_paragraph(text, style)
        anchor = self.doc.paragraphs[index]
        new_p = anchor.insert_paragraph_before(text)
        warn = self._apply_style(new_p, style)
        self.save()
        return f"Inserted paragraph at index {index}." + (f" {warn}" if warn else "")

    def set_paragraph(self, index: int, text: str, style: str | None = None) -> str:
        n = len(self.doc.paragraphs)
        if index < 0 or index >= n:
            return f"Error: index {index} out of range (0..{n - 1})."
        p = self.doc.paragraphs[index]
        self._clear_runs(p)
        p.add_run(text)
        warn = self._apply_style(p, style)
        self.save()
        return f"Updated paragraph at index {index}." + (f" {warn}" if warn else "")

    def delete_paragraph(self, index: int) -> str:
        n = len(self.doc.paragraphs)
        if index < 0 or index >= n:
            return f"Error: index {index} out of range (0..{n - 1})."
        p = self.doc.paragraphs[index]
        p._element.getparent().remove(p._element)
        self.save()
        return f"Deleted paragraph at index {index}."

    @staticmethod
    def _replace_in_para(p, find: str, replace: str) -> int:
        """Replace text within a paragraph while preserving run-level formatting.

        Strategy:
        1. Try run-level replacement first — this preserves bold, italic, color,
           font size, etc. on every run that contains the target text entirely
           within one run.
        2. Fall back to a whole-paragraph rebuild only when the target string
           spans multiple runs (rare, but possible when Word splits text
           arbitrarily). In that case formatting IS lost, which is the same
           behaviour as before — still better than silently missing the match.
        """
        if find not in p.text:
            return 0

        # --- pass 1: run-level (format-preserving) ---
        count = 0
        for run in p.runs:
            if find in run.text:
                count += run.text.count(find)
                run.text = run.text.replace(find, replace)

        # --- pass 2: cross-run fallback (format-lossy, but correct) ---
        # Re-read p.text after pass 1; if find still appears it crosses run
        # boundaries and we have no choice but to rebuild the paragraph.
        if find in p.text:
            remaining = p.text.count(find)
            count += remaining
            new_text = p.text.replace(find, replace)
            for run in list(p.runs):
                run._element.getparent().remove(run._element)
            p.add_run(new_text)

        return count

    def replace_text(self, find: str, replace: str) -> str:
        if not find:
            return "Error: 'find' must be non-empty."
        count = 0

        for kind, obj in self._body_blocks():
            if kind == "paragraph":
                count += self._replace_in_para(obj, find, replace)
            else:
                for row in obj.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            count += self._replace_in_para(p, find, replace)

        if count:
            self.save()
        return f"Replaced {count} occurrence(s) of {find!r}."

    def read_table(self, table_index: int) -> str:
        """Return detailed cell-by-cell content of a specific table."""
        tables = self.doc.tables
        if not tables:
            return "This document contains no tables."
        if table_index < 0 or table_index >= len(tables):
            return f"Error: table_index {table_index} out of range (0..{len(tables) - 1})."
        table = tables[table_index]
        lines = [f"Table {table_index}: {len(table.rows)} rows × {len(table.columns)} cols"]
        for r, row in enumerate(table.rows):
            for c, cell in enumerate(row.cells):
                cell_text = cell.text.replace("\n", " / ")
                lines.append(f"  [{r}][{c}]: {cell_text!r}")
        return "\n".join(lines)

    def set_table_cell(self, table_index: int, row: int, col: int, text: str) -> str:
        """Replace the text of one table cell, preserving the cell's paragraph style."""
        tables = self.doc.tables
        if not tables:
            return "This document contains no tables."
        if table_index < 0 or table_index >= len(tables):
            return f"Error: table_index {table_index} out of range (0..{len(tables) - 1})."
        table = tables[table_index]
        if row < 0 or row >= len(table.rows):
            return f"Error: row {row} out of range (0..{len(table.rows) - 1})."
        row_cells = table.rows[row].cells
        if col < 0 or col >= len(row_cells):
            return f"Error: col {col} out of range (0..{len(row_cells) - 1})."
        cell = row_cells[col]
        paras = cell.paragraphs
        if paras:
            p = paras[0]
            # Preserve the first run's character formatting (bold, font, size…)
            # by writing into its text property rather than clearing all runs.
            if p.runs:
                p.runs[0].text = text
                # Remove any additional runs so there's no leftover text.
                for extra_run in p.runs[1:]:
                    extra_run._element.getparent().remove(extra_run._element)
            else:
                p.add_run(text)
            # Remove any extra paragraphs inside the cell.
            for extra in paras[1:]:
                extra._element.getparent().remove(extra._element)
        else:
            cell.add_paragraph(text)
        self.save()
        return f"Updated table[{table_index}] row {row} col {col} → {text!r}."


# ---------- HTML round-trip (used by the manual editor in the Streamlit UI) ----------

# Tags Quill emits → Word style names. Anything else falls through to "Normal".
_TAG_TO_STYLE = {
    "h1": "Heading 1",
    "h2": "Heading 2",
    "h3": "Heading 3",
    "blockquote": "Quote",
    "p": "Normal",
}


def doc_to_html(editor: "DocEditor") -> str:
    """Render the document as Quill-compatible HTML (block + simple inline tags)."""
    parts: list[str] = []
    list_kind: str | None = None  # "ul" | "ol" | None

    def close_list() -> None:
        nonlocal list_kind
        if list_kind:
            parts.append(f"</{list_kind}>")
            list_kind = None

    for p in editor.doc.paragraphs:
        style = p.style.name if p.style else "Normal"

        # Build the inner HTML from runs so bold/italic/underline survive a round-trip.
        inner_parts: list[str] = []
        for run in p.runs:
            text = escape(run.text or "")
            if not text:
                continue
            if run.bold:
                text = f"<strong>{text}</strong>"
            if run.italic:
                text = f"<em>{text}</em>"
            if run.underline:
                text = f"<u>{text}</u>"
            inner_parts.append(text)
        inner = "".join(inner_parts) or escape(p.text or "")
        if not inner:
            inner = "<br>"

        if style == "List Bullet":
            if list_kind != "ul":
                close_list()
                parts.append("<ul>")
                list_kind = "ul"
            parts.append(f"<li>{inner}</li>")
            continue
        if style == "List Number":
            if list_kind != "ol":
                close_list()
                parts.append("<ol>")
                list_kind = "ol"
            parts.append(f"<li>{inner}</li>")
            continue

        close_list()
        if style == "Title":
            # Quill has no "title" — represent with H1; we'll round-trip back as H1/Title best-effort.
            parts.append(f"<h1 data-style=\"Title\">{inner}</h1>")
        elif style == "Heading 1":
            parts.append(f"<h1>{inner}</h1>")
        elif style == "Heading 2":
            parts.append(f"<h2>{inner}</h2>")
        elif style == "Heading 3":
            parts.append(f"<h3>{inner}</h3>")
        elif style == "Quote":
            parts.append(f"<blockquote>{inner}</blockquote>")
        else:
            parts.append(f"<p>{inner}</p>")

    close_list()
    return "".join(parts) or "<p><br></p>"


def html_to_doc(editor: "DocEditor", html: str) -> None:
    """Parse Quill HTML and replace the document body. Saves on success.

    Imports BeautifulSoup lazily so the CLI editor doesn't require it.
    """
    from bs4 import BeautifulSoup, NavigableString

    soup = BeautifulSoup(html or "", "html.parser")
    body = soup.body or soup

    # Wipe existing paragraphs.
    for p in list(editor.doc.paragraphs):
        p._element.getparent().remove(p._element)

    def add_block(text_runs: list[tuple[str, dict]], style: str) -> None:
        """Append one paragraph from a list of (text, formatting) runs."""
        p = editor.doc.add_paragraph()
        editor._apply_style(p, style)
        if not text_runs:
            return
        for text, fmt in text_runs:
            if not text:
                continue
            run = p.add_run(text)
            if fmt.get("bold"):
                run.bold = True
            if fmt.get("italic"):
                run.italic = True
            if fmt.get("underline"):
                run.underline = True

    def collect_runs(node, fmt: dict) -> list[tuple[str, dict]]:
        """Walk inline children of a block element, returning (text, fmt) runs."""
        runs: list[tuple[str, dict]] = []
        for child in node.children:
            if isinstance(child, NavigableString):
                txt = str(child)
                if txt:
                    runs.append((txt, dict(fmt)))
                continue
            tag = child.name.lower()
            child_fmt = dict(fmt)
            if tag in ("strong", "b"):
                child_fmt["bold"] = True
            elif tag in ("em", "i"):
                child_fmt["italic"] = True
            elif tag == "u":
                child_fmt["underline"] = True
            elif tag == "br":
                runs.append(("\n", dict(fmt)))
                continue
            runs.extend(collect_runs(child, child_fmt))
        return runs

    # Walk top-level blocks.
    for el in body.children:
        if isinstance(el, NavigableString):
            txt = str(el).strip()
            if txt:
                add_block([(txt, {})], "Normal")
            continue

        tag = el.name.lower() if el.name else ""
        if tag in ("ul", "ol"):
            style = "List Bullet" if tag == "ul" else "List Number"
            for li in el.find_all("li", recursive=False):
                add_block(collect_runs(li, {}), style)
            continue

        if tag == "h1" and el.get("data-style") == "Title":
            style = "Title"
        else:
            style = _TAG_TO_STYLE.get(tag, "Normal")
        add_block(collect_runs(el, {}), style)

    # Always keep at least one paragraph so python-docx stays happy.
    if not editor.doc.paragraphs:
        editor.doc.add_paragraph("")

    editor.save()


def dispatch(editor: DocEditor, name: str, args: dict) -> str:
    method = getattr(editor, name, None)
    if method is None or name.startswith("_"):
        return f"Unknown tool: {name}"
    try:
        return method(**args)
    except TypeError as e:
        return f"Bad arguments for {name}: {e}"
    except Exception as e:  # noqa: BLE001
        return f"Error in {name}: {e}"


def run_turn(client: Anthropic, editor: DocEditor, history: list[dict]) -> None:
    while True:
        print("Assistant: ", end="", flush=True)
        with client.messages.stream(
            model=MODEL,
            max_tokens=16000,
            system=SYSTEM_PROMPT,
            messages=history,
            tools=TOOLS,
            thinking={"type": "adaptive"},
        ) as stream:
            for event in stream:
                if (
                    event.type == "content_block_delta"
                    and getattr(event.delta, "type", None) == "text_delta"
                ):
                    print(event.delta.text, end="", flush=True)
            final = stream.get_final_message()
        print()

        history.append({"role": "assistant", "content": final.content})

        if final.stop_reason != "tool_use":
            print()
            return

        tool_results = []
        for block in final.content:
            if block.type != "tool_use":
                continue
            preview = json.dumps(block.input)
            if len(preview) > 140:
                preview = preview[:137] + "..."
            print(f"  -> {block.name}({preview})")
            result = dispatch(editor, block.name, block.input)
            tool_results.append(
                {
                    "type": "tool_result",
                    "tool_use_id": block.id,
                    "content": result,
                }
            )
        history.append({"role": "user", "content": tool_results})


def chat(path: Path) -> None:
    if not os.environ.get("ANTHROPIC_API_KEY"):
        print("ERROR: ANTHROPIC_API_KEY is not set.")
        print("Copy .env.example to .env and fill in your key.")
        sys.exit(1)

    editor = DocEditor(path)
    client = Anthropic()
    history: list[dict] = []

    print(f"Word agent ready. Editing: {path.resolve()}")
    print(f"Model: {MODEL}. Type 'exit' to leave, 'reset' to clear history.\n")

    while True:
        try:
            user_input = input("You: ").strip()
        except (EOFError, KeyboardInterrupt):
            print("\nGoodbye.")
            return

        if not user_input:
            continue
        if user_input.lower() in {"exit", "quit"}:
            print("Goodbye.")
            return
        if user_input.lower() == "reset":
            history = []
            print("[Conversation history cleared]\n")
            continue

        history.append({"role": "user", "content": user_input})

        try:
            run_turn(client, editor, history)
        except RateLimitError:
            print("\n[Rate limited - wait a moment and try again]\n")
            history.pop()
        except APIConnectionError as e:
            print(f"\n[Connection error: {e}]\n")
            history.pop()
        except APIError as e:
            print(f"\n[API error: {e}]\n")
            history.pop()


def main() -> None:
    if len(sys.argv) < 2:
        print("Usage: python doc_editor.py <path-to-document.docx>")
        print("If the file doesn't exist, a new empty document will be created.")
        sys.exit(1)
    path = Path(sys.argv[1])
    if path.suffix.lower() != ".docx":
        print("Error: file must have a .docx extension.")
        sys.exit(1)
    chat(path)


if __name__ == "__main__":
    main()
